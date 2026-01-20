"""Excel 模板填入模組"""

from pathlib import Path

from openpyxl import load_workbook


def write_summary(
    template_path: Path,
    output_path: Path,
    quantity: int,  # 支數 → B3
    sheet_count: int,  # 張數 → C3
    purchase_order_no: str,  # 採購單號 → G2
) -> None:
    """
    讀取 Excel 模板，填入支數和張數，另存新檔

    Args:
        template_path: 模板檔案路徑 (template.xlsx)
        output_path: 輸出檔案路徑 ({採購單號}-單據明細.xlsx)
        quantity: 支數，填入 B3
        sheet_count: 張數，填入 C3
        purchase_order_no: 採購單號，填入 G2
    """
    # 讀取模板（不修改原檔）
    wb = load_workbook(template_path)
    ws = wb.active

    # 填入數值
    ws["B3"] = quantity  # 支數
    ws["C3"] = sheet_count  # 張數
    ws["G2"] = purchase_order_no  # 採購單號

    # 另存新檔
    wb.save(output_path)
    wb.close()
