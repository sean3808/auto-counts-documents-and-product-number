"""phase2 模組測試"""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from doc_processor.logger import ProcessLogger
from doc_processor.phase1 import run_phase1
from doc_processor.phase2 import run_phase2


class TestRunPhase2:
    """Phase 2 整合測試"""

    def test_generate_excel(self, fixtures_dir: Path, temp_output_dir: Path):
        """測試 Excel 生成"""
        template_path = fixtures_dir / "template.xlsx"

        # 先執行 Phase 1 產生合併 PDF
        logger = ProcessLogger(temp_output_dir)
        run_phase1(fixtures_dir, temp_output_dir, logger)

        # 執行 Phase 2
        logger2 = ProcessLogger(temp_output_dir)
        exit_code = run_phase2(temp_output_dir, template_path, logger2)

        # 應該成功
        assert exit_code == 0

        # 檢查 Excel 輸出
        excel_files = list(temp_output_dir.glob("*-單據明細.xlsx"))
        assert len(excel_files) == 1

        # 驗證內容
        wb = load_workbook(excel_files[0])
        ws = wb.active
        assert ws["B3"].value == 1  # 支數
        assert ws["C3"].value == 4  # 張數
        wb.close()

    def test_no_pdf_files(self, tmp_path: Path, fixtures_dir: Path):
        """測試沒有 PDF 檔案的情況"""
        empty_output = tmp_path / "empty_output"
        empty_output.mkdir()
        template_path = fixtures_dir / "template.xlsx"

        logger = ProcessLogger(empty_output)
        exit_code = run_phase2(empty_output, template_path, logger)

        # 沒有檔案應該回傳成功
        assert exit_code == 0
