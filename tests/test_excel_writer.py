"""excel_writer 模組測試"""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from doc_processor.excel_writer import write_summary


class TestWriteSummary:
    """Excel 寫入測試"""

    def test_write_values(self, fixtures_dir: Path, tmp_path: Path):
        """測試寫入支數和張數"""
        template_path = fixtures_dir / "template.xlsx"
        if not template_path.exists():
            pytest.skip("模板檔案不存在")

        output_path = tmp_path / "test_output.xlsx"

        write_summary(template_path, output_path, quantity=5, sheet_count=10, purchase_order_no="1234567890123")

        # 驗證輸出
        wb = load_workbook(output_path)
        ws = wb.active
        assert ws["B3"].value == 5  # 支數
        assert ws["C3"].value == 10  # 張數
        assert ws["G2"].value == "1234567890123"  # 採購單號
        wb.close()

    def test_template_not_modified(self, fixtures_dir: Path, tmp_path: Path):
        """測試模板檔案不被修改"""
        template_path = fixtures_dir / "template.xlsx"
        if not template_path.exists():
            pytest.skip("模板檔案不存在")

        # 讀取模板原始值
        wb = load_workbook(template_path)
        ws = wb.active
        original_b3 = ws["B3"].value
        original_c3 = ws["C3"].value
        wb.close()

        # 執行寫入
        output_path = tmp_path / "test_output.xlsx"
        write_summary(template_path, output_path, quantity=99, sheet_count=99, purchase_order_no="9999999999999")

        # 確認模板未被修改
        wb = load_workbook(template_path)
        ws = wb.active
        assert ws["B3"].value == original_b3
        assert ws["C3"].value == original_c3
        wb.close()
